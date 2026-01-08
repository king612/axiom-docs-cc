"""
Excel parser service for extracting project data from bid spreadsheets.

Spreadsheet structure (first tab):
- Rows 1-12: Project info (name, proposal number, sqft)
- Rows 58-90: Fee estimates and totals
- Rows 94+: "Identify (3) past projects to compare" training data
- Rows 106+: "Recent Proposals That are similar size and scope" training data
"""

import logging
import re
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from backend.models.project import Project, ProjectFeatures

logger = logging.getLogger(__name__)

# Column mappings for "past projects to compare" table (row 95 headers)
TRAINING_TABLE_COLUMNS = {
    "job_id": 1,           # Column A
    "job_name": 3,         # Column C
    "total_sheets": 5,     # Column E
    "engineering_hrs": 6,  # Column F
    "drafting_hrs": 7,     # Column G
    "manager_hrs": 8,      # Column H
    "cd_fee": 14,          # Column N (Design Fee)
    "ca_fee": 15,          # Column O
    "sq_ft": 16,           # Column P
    "total_spent_cd": 17,  # Column Q (Total Dollars spent Design)
    "total_spent_ca": 18,  # Column R
}

# Column mappings for "Recent Proposals" table
RECENT_PROPOSALS_COLUMNS = {
    "proposal_number": 1,  # Column A
    "proposal_name": 2,    # Column B
    "sq_ft": 3,            # Column C
    "cd_fee": 5,           # Column E (Proposal Design Fee)
    "ca_fee": 6,           # Column F (Proposal CA Fee)
}


def safe_float(value) -> Optional[float]:
    """Safely convert a value to float, returning None if not possible."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        # Handle currency formatting
        cleaned = value.replace("$", "").replace(",", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def safe_int(value) -> Optional[int]:
    """Safely convert a value to int, returning None if not possible."""
    f = safe_float(value)
    if f is not None:
        return int(f)
    return None


def extract_project_id_from_filename(filename: str) -> str:
    """
    Extract project ID from filename.

    Examples:
        "P22-003B ABS Tracy.xlsx" -> "P22-003B"
        "P24-310A Sonora.xlsx" -> "P24-310A"
    """
    match = re.match(r"^(P\d{2}-\d{3}[A-Z]?)", filename)
    if match:
        return match.group(1)
    # Fallback: return filename without extension
    return Path(filename).stem


def extract_year_from_path(filepath: Path) -> int:
    """Extract year from file path (e.g., data/estimates/2024/...)."""
    for part in filepath.parts:
        if part.isdigit() and len(part) == 4:
            return int(part)
    return 0


def find_section_row(ws: Worksheet, search_text: str, start_row: int = 1, end_row: int = 150) -> Optional[int]:
    """Find the row containing a specific text in column A."""
    search_lower = search_text.lower()
    for row in range(start_row, end_row + 1):
        cell_value = ws.cell(row=row, column=1).value
        if cell_value and search_lower in str(cell_value).lower():
            return row
    return None


def extract_training_projects(ws: Worksheet) -> list[str]:
    """
    Extract project IDs from training data tables.

    Returns list of project IDs that domain experts selected as similar.
    """
    similar_projects = []

    # Find "past projects to compare" section
    past_projects_row = find_section_row(ws, "past projects to compare")
    if past_projects_row:
        # Data starts 2 rows after header (row 95 is headers, 96+ is data)
        data_start = past_projects_row + 2
        for row in range(data_start, data_start + 15):  # Check up to 15 rows
            job_id = ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["job_id"]).value
            if job_id and isinstance(job_id, str) and job_id.strip():
                # Normalize project ID (A23-102 style)
                similar_projects.append(job_id.strip())
            else:
                # Stop if we hit an empty row
                break

    # Find "Recent Proposals" section
    recent_row = find_section_row(ws, "Recent Proposals That are similar")
    if recent_row:
        data_start = recent_row + 2  # Skip header row
        for row in range(data_start, data_start + 10):
            proposal_num = ws.cell(row=row, column=RECENT_PROPOSALS_COLUMNS["proposal_number"]).value
            if proposal_num and isinstance(proposal_num, str) and proposal_num.strip():
                similar_projects.append(proposal_num.strip())
            else:
                break

    return similar_projects


def extract_current_project_features(ws: Worksheet) -> ProjectFeatures:
    """
    Extract features for the current project from the spreadsheet.

    Looks for totals in the fee estimate section.
    """
    features = ProjectFeatures()

    # Get SQ FT from row 8, column 2
    features.sq_ft = safe_int(ws.cell(row=8, column=2).value)

    # Find totals section
    # Look for "CDs Total" row for design totals
    cd_total_row = find_section_row(ws, "CDs Total", start_row=60, end_row=90)
    if cd_total_row:
        # These are typically in specific columns based on the fee table structure
        pass  # Structure varies, would need more samples to determine

    # Look for "Grand Totals" row
    grand_total_row = find_section_row(ws, "Grand Totals", start_row=80, end_row=95)
    if grand_total_row:
        # Column 11 often has total hours
        total_hrs = safe_float(ws.cell(row=grand_total_row, column=11).value)
        # Column 17 often has total fee
        total_fee = safe_float(ws.cell(row=grand_total_row, column=17).value)

    # Look for specific hour breakdowns
    # Find "Total Hours (All Stages)" for each section
    for row in range(15, 60):
        cell_val = ws.cell(row=row, column=1).value
        if cell_val and "Total Hours (All Stages)" in str(cell_val):
            hours = safe_float(ws.cell(row=row, column=2).value)
            # Determine which section this is based on preceding rows
            for prev_row in range(row - 5, row):
                prev_val = ws.cell(row=prev_row, column=1).value
                if prev_val:
                    prev_str = str(prev_val).lower()
                    if "drafting" in prev_str and features.drafting_hrs is None:
                        features.drafting_hrs = hours
                    elif "engineering" in prev_str and features.engineering_hrs is None:
                        features.engineering_hrs = hours
                    elif "management" in prev_str and features.manager_hrs is None:
                        features.manager_hrs = hours
                    break

    # Get fees from the fee estimate section
    cd_fee_row = find_section_row(ws, "CDs Total", start_row=65, end_row=85)
    if cd_fee_row:
        features.cd_fee = safe_float(ws.cell(row=cd_fee_row, column=17).value)

    ca_fee_row = find_section_row(ws, "CA Total", start_row=80, end_row=90)
    if ca_fee_row:
        features.ca_fee = safe_float(ws.cell(row=ca_fee_row, column=17).value)

    # Count sheets from sheet count section
    totals_row = find_section_row(ws, "Totals", start_row=25, end_row=35)
    if totals_row:
        features.total_sheets = safe_int(ws.cell(row=totals_row, column=3).value)

    return features


def extract_training_table_features(ws: Worksheet) -> list[tuple[str, ProjectFeatures]]:
    """
    Extract features from the training data table for each referenced project.

    Returns list of (project_id, features) tuples.
    """
    results = []

    past_projects_row = find_section_row(ws, "past projects to compare")
    if not past_projects_row:
        return results

    # Data starts 2 rows after the section header
    data_start = past_projects_row + 2

    for row in range(data_start, data_start + 15):
        job_id = ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["job_id"]).value
        if not job_id or not isinstance(job_id, str) or not job_id.strip():
            break

        features = ProjectFeatures(
            total_sheets=safe_int(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["total_sheets"]).value),
            engineering_hrs=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["engineering_hrs"]).value),
            drafting_hrs=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["drafting_hrs"]).value),
            manager_hrs=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["manager_hrs"]).value),
            cd_fee=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["cd_fee"]).value),
            ca_fee=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["ca_fee"]).value),
            sq_ft=safe_int(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["sq_ft"]).value),
            total_spent_cd=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["total_spent_cd"]).value),
            total_spent_ca=safe_float(ws.cell(row=row, column=TRAINING_TABLE_COLUMNS["total_spent_ca"]).value),
        )

        results.append((job_id.strip(), features))

    return results


def parse_excel_file(filepath: Path) -> Optional[Project]:
    """
    Parse an Excel bid spreadsheet and extract project data.

    Args:
        filepath: Path to the Excel file

    Returns:
        Project object with extracted data, or None if parsing fails
    """
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.worksheets[0]  # Always use first sheet

        # Extract project ID from filename
        project_id = extract_project_id_from_filename(filepath.name)

        # Extract year from path
        year = extract_year_from_path(filepath)

        # Extract features for the current project
        features = extract_current_project_features(ws)

        # Extract training labels (similar projects)
        similar_projects = extract_training_projects(ws)

        project = Project(
            project_id=project_id,
            filename=filepath.name,
            filepath=filepath,
            year=year,
            features=features,
            similar_projects=similar_projects,
            has_training_labels=len(similar_projects) > 0,
        )

        wb.close()
        return project

    except Exception as e:
        logger.error(f"Error parsing {filepath}: {e}")
        return None


def parse_training_data(filepath: Path) -> list[tuple[str, ProjectFeatures]]:
    """
    Parse training data from an Excel file.

    Returns list of (project_id, features) from the "past projects to compare" table.
    """
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.worksheets[0]
        results = extract_training_table_features(ws)
        wb.close()
        return results
    except Exception as e:
        logger.error(f"Error parsing training data from {filepath}: {e}")
        return []
